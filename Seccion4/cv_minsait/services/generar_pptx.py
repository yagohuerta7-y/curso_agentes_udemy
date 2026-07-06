import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from typing import Optional


def generar_pptx(data_cv: dict, plantilla_path: str, salida_path: str):
    shutil.copy(plantilla_path, salida_path)

    prs = Presentation(salida_path)
    slide = prs.slides[0]

    _reemplazar_y_ocultar(slide, data_cv)

    prs.save(salida_path)
    return salida_path


def _reemplazar_y_ocultar(slide, data_cv: dict):
    shapes = list(slide.shapes)

    _procesar_educacion(shapes[1], data_cv)
    _procesar_contenido_principal(shapes[6], data_cv)

    for shape in shapes:
        if shape.has_text_frame:
            _reemplazar_placeholders_en_shape(shape, data_cv)


def _reemplazar_placeholders_en_shape(shape, data_cv: dict):
    if not shape.has_text_frame:
        return

    map_reemplazos = _obtener_mapa_reemplazos(data_cv)

    for paragraph in shape.text_frame.paragraphs:
        texto_completo = paragraph.text
        if not any(ph in texto_completo for ph in map_reemplazos.keys()):
            continue

        texto_nuevo = texto_completo
        for placeholder, valor in map_reemplazos.items():
            texto_nuevo = texto_nuevo.replace(placeholder, valor)

        if texto_nuevo != texto_completo:
            _reemplazar_texto_parrafo(paragraph, texto_nuevo)


def _obtener_mapa_reemplazos(data_cv: dict) -> dict:
    reemplazos = {
        "{{nombre}}": data_cv.get("nombre", ""),
        "{{licenciatura}}": data_cv.get("licenciatura", ""),
        "{{maestria}}": data_cv.get("maestria") or "",
        "{{hab_duras}}": data_cv.get("hab_duras", ""),
        "{{herramientas}}": data_cv.get("herramientas", ""),
        "{{idiomas}}": data_cv.get("idiomas", ""),
        "{{exp_prof}}": data_cv.get("exp_prof", ""),
    }

    cursos = data_cv.get("cursos", [])
    reemplazos["{{cursos}}"] = "\n".join(cursos) if isinstance(cursos, list) else str(cursos)

    trabajos = data_cv.get("trabajos", [])
    for i in range(1, 4):
        if i <= len(trabajos):
            trabajo = trabajos[i - 1]
            reemplazos[f"{{{{empresa_{i}}}}}"] = trabajo.get("empresa", "")
            reemplazos[f"{{{{puesto_{i}}}}}"] = trabajo.get("puesto", "")
            reemplazos[f"{{{{fecha_inicio_{i}}}}}"] = trabajo.get("fecha_inicio", "")
            fecha_fin = trabajo.get("fecha_fin")
            reemplazos[f"{{{{fecha_fin_{i}}}}}"] = fecha_fin if fecha_fin else "Actual"
            actividades = trabajo.get("actividades", [])
            if len(actividades) >= 3:
                reemplazos[f"{{{{actividad1_{i}}}}}"] = actividades[0]
                reemplazos[f"{{{{actividad2_{i}}}}}"] = actividades[1]
                reemplazos[f"{{{{actividad3_{i}}}}}"] = actividades[2]
        else:
            reemplazos[f"{{{{empresa_{i}}}}}"] = ""
            reemplazos[f"{{{{puesto_{i}}}}}"] = ""
            reemplazos[f"{{{{fecha_inicio_{i}}}}}"] = ""
            reemplazos[f"{{{{fecha_fin_{i}}}}}"] = ""
            reemplazos[f"{{{{actividad1_{i}}}}}"] = ""
            reemplazos[f"{{{{actividad2_{i}}}}}"] = ""
            reemplazos[f"{{{{actividad3_{i}}}}}"] = ""

    return reemplazos


def _reemplazar_texto_parrafo(paragraph, texto_nuevo: str):
    if not paragraph.runs:
        return

    primer_run = paragraph.runs[0]
    primer_run.text = texto_nuevo

    for run in paragraph.runs[1:]:
        run.text = ""


def _procesar_educacion(shape, data_cv: dict):
    if not shape.has_text_frame:
        return

    if not data_cv.get("maestria"):
        _eliminar_parrafo_por_placeholder(shape, "{{maestria}}")

    _reemplazar_placeholders_en_shape(shape, data_cv)


def _procesar_contenido_principal(shape, data_cv: dict):
    if not shape.has_text_frame:
        return

    trabajos = data_cv.get("trabajos", [])

    for n in [2, 3]:
        if n > len(trabajos):
            _ocultar_seccion_trabajo(shape, n)

    _reemplazar_placeholders_en_shape(shape, data_cv)


def _eliminar_parrafo_por_placeholder(shape, placeholder: str):
    paragraphs = list(shape.text_frame.paragraphs)
    for p in paragraphs:
        if placeholder in p.text:
            p._p.getparent().remove(p._p)
            return


def _ocultar_seccion_trabajo(shape, n: int):
    placeholders_seccion = [
        f"empresa_{n}",
        f"puesto_{n}",
        f"fecha_inicio_{n}",
        f"fecha_fin_{n}",
        f"actividad1_{n}",
        f"actividad2_{n}",
        f"actividad3_{n}",
    ]

    paragraphs_elements = list(shape.text_frame._txBody.findall(
        '{http://schemas.openxmlformats.org/drawingml/2006/main}p'
    ))

    paragraphs_a_eliminar = []
    for p_elem in paragraphs_elements:
        texto = "".join(t.text or "" for t in p_elem.findall(
            './/{http://schemas.openxmlformats.org/drawingml/2006/main}t'
        ))
        for ph in placeholders_seccion:
            if f"{{{{{ph}}}}}" in texto:
                paragraphs_a_eliminar.append(p_elem)
                break

    if paragraphs_a_eliminar:
        primer_p = paragraphs_a_eliminar[0]
        idx = paragraphs_elements.index(primer_p)
        if idx > 0:
            parrafo_anterior = paragraphs_elements[idx - 1]
            texto_anterior = "".join(t.text or "" for t in parrafo_anterior.findall(
                './/{http://schemas.openxmlformats.org/drawingml/2006/main}t'
            ))
            if not texto_anterior.strip():
                paragraphs_a_eliminar.append(parrafo_anterior)

    for p_elem in paragraphs_a_eliminar:
        try:
            p_elem.getparent().remove(p_elem)
        except Exception:
            pass
